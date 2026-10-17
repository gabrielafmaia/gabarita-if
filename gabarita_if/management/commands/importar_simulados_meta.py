
import re
import unicodedata
from pathlib import Path
from html import escape

from docx import Document
from docx.oxml.ns import qn
from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from gabarita_if.models import (
    Assunto,
    Avaliacao,
    Disciplina,
    Fonte,
    Questao,
    TextoApoio,
)


PASTA_PADRAO = (
    Path.home()
    / "Downloads"
    / "simulados_importacao"
    / "Simulados Meta IFRN"
)

PADRAO_ARQUIVO = re.compile(
    r"^(I{1,2})\s+Simulad[aã]o\s+(20\d{2})\s*-\s*(.+)\.docx$",
    re.I,
)

PADRAO_QUESTAO = re.compile(
    r"^\s*(\d{1,2})[.)]\s*(.*)$"
)

PADRAO_ALTERNATIVA = re.compile(
    r"^\s*([A-Da-d])[).]\s*(.*?)\s*$"
)

PADRAO_GABARITO = re.compile(
    r"^\s*(\d{1,2})\s*[-:.)]\s*([A-Da-d])\s*$"
)

PADRAO_TEXTO = re.compile(
    r"^TEXTO\s+([IVX]+|\d+)\s*$",
    re.I,
)

SECOES = {
    "PORTUGUES": "Português",
    "MATEMATICA": "Matemática",
    "CIENCIAS": "Ciências",
    "HISTORIA": "História",
    "GEOGRAFIA": "Geografia",
    "INGLES": "Inglês",
}


def normalizar(texto):
    texto = unicodedata.normalize("NFD", texto)

    return "".join(
        caractere
        for caractere in texto
        if unicodedata.category(caractere) != "Mn"
    ).upper().strip()


def ler_paragrafos(documento):
    linhas = [
        paragrafo.text.strip()
        for paragrafo in documento.paragraphs
        if paragrafo.text.strip()
    ]

    for tabela in documento.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                linhas.extend(
                    paragrafo.text.strip()
                    for paragrafo in celula.paragraphs
                    if paragrafo.text.strip()
                )

    return linhas


def extrair_gabarito(linhas):
    respostas = {}
    em_gabarito = False
    letras = []

    for linha in linhas:
        texto = linha.strip()
        chave = normalizar(texto).rstrip(" :")

        if chave in {
            "GABARITO",
            "GABARITO OFICIAL",
            "RESPOSTAS",
        }:
            em_gabarito = True
            continue

        if not em_gabarito:
            continue

        correspondencia = PADRAO_GABARITO.fullmatch(texto)

        if correspondencia:
            numero = int(correspondencia.group(1))
            letra = correspondencia.group(2).upper()
            respostas[numero] = letra

        elif re.fullmatch(r"[A-Da-d]", texto):
            letras.append(texto.upper())

        else:
            pares = re.findall(
                r"(\d{1,2})\s*[-:.)]\s*([A-Da-d])",
                texto,
            )

            for numero, letra in pares:
                respostas[int(numero)] = letra.upper()

    for numero, letra in enumerate(letras[:40], 1):
        respostas.setdefault(numero, letra)

    return respostas


def obter_numeracao(paragrafo):
    propriedades = paragrafo._p.pPr

    if (
        propriedades is None
        or propriedades.numPr is None
        or propriedades.numPr.numId is None
    ):
        return None, None

    numeracao = propriedades.numPr

    nivel = (
        numeracao.ilvl.val
        if numeracao.ilvl is not None
        else 0
    )

    return numeracao.numId.val, nivel


def imagem_ids(paragrafo):
    return [
        elemento.get(qn("r:embed"))
        for elemento in paragrafo._p.iter(qn("a:blip"))
        if elemento.get(qn("r:embed"))
    ]


def html_paragrafo(
    paragrafo,
    documento,
    salvar=False,
    prefixo="",
):
    texto = paragrafo.text.strip()
    partes = []

    if texto:
        partes.append(
            f"<p>{escape(texto)}</p>"
        )

    for rid in imagem_ids(paragrafo):
        parte = documento.part.related_parts.get(rid)

        if parte is None:
            continue

        extensao = Path(
            str(parte.partname)
        ).suffix.lower()

        if extensao not in {
            ".png",
            ".jpg",
            ".jpeg",
            ".gif",
            ".webp",
            ".svg",
        }:
            extensao = ".png"

        nome = f"{prefixo}_{rid}{extensao}"

        caminho = (
            Path(settings.MEDIA_ROOT)
            / "simulados_meta"
            / nome
        )

        if salvar:
            caminho.parent.mkdir(
                parents=True,
                exist_ok=True,
            )
            caminho.write_bytes(parte.blob)

        url = (
            str(settings.MEDIA_URL).rstrip("/")
            + "/simulados_meta/"
            + nome
        )

        partes.append(
            f'<p><img src="{escape(url, quote=True)}" '
            f'alt="Imagem do simulado" '
            f'style="max-width:100%;height:auto"></p>'
        )

    return "\n".join(partes)


def chave_texto(valor):
    valor = normalizar(valor)

    numeros = {
        "I": 1,
        "II": 2,
        "III": 3,
        "IV": 4,
        "V": 5,
        "VI": 6,
        "VII": 7,
        "VIII": 8,
        "IX": 9,
        "X": 10,
    }

    if valor in numeros:
        return numeros[valor]

    if valor.isdigit():
        return int(valor)

    return None


def extrair_apoios(
    documento,
    limite,
    salvar=False,
    prefixo="",
):
    apoios = []
    atual = None

    for paragrafo in documento.paragraphs[:limite]:
        texto = paragrafo.text.strip()

        correspondencia = PADRAO_TEXTO.fullmatch(
            normalizar(texto)
        )

        if correspondencia:
            atual = {
                "numero": chave_texto(
                    correspondencia.group(1)
                ),
                "titulo": texto,
                "partes": [],
            }

            apoios.append(atual)
            continue

        if atual is not None:
            conteudo = html_paragrafo(
                paragrafo,
                documento,
                salvar,
                prefixo,
            )

            if conteudo:
                atual["partes"].append(conteudo)

    return [
        {
            "numero": apoio["numero"],
            "titulo": apoio["titulo"],
            "texto": "\n".join(apoio["partes"]),
        }
        for apoio in apoios
    ]


def textos_referidos(enunciado):
    numeros = set()

    referencias = re.findall(
        r"\bTEXTOS?\s+([IVX]+|\d+)\b",
        normalizar(enunciado),
    )

    for referencia in referencias:
        numero = chave_texto(referencia)

        if numero is not None:
            numeros.add(numero)

    return numeros


def html_intervalo(
    paragrafos,
    inicio,
    fim,
    documento,
    salvar,
    prefixo,
):
    partes = []

    for indice in range(inicio, fim):
        paragrafo = paragrafos[indice]

        if (
            paragrafo.text.strip()
            or imagem_ids(paragrafo)
        ):
            conteudo = html_paragrafo(
                paragrafo,
                documento,
                salvar,
                prefixo,
            )

            if conteudo:
                partes.append(conteudo)

    return "\n".join(partes)


def eh_resolucao(texto):
    chave = normalizar(texto).strip().rstrip(" :.")

    return chave in {
        "RESOLUCAO",
        "RESPOSTA",
        "RESPOSTA COMENTADA",
        "COMENTARIO",
        "JUSTIFICATIVA",
        "EXPLICACAO",
        "VER RESPOSTA",
    }


def extrair_questoes_2020(
    documento,
    gabaritos,
    titulo,
    tema,
    ano,
    salvar=False,
    prefixo="",
):
    paragrafos = documento.paragraphs
    inicios = []
    disciplinas = {}
    disciplina = "Português"

    for indice, paragrafo in enumerate(paragrafos):
        chave = normalizar(
            paragrafo.text
        ).rstrip(" :")

        if chave in SECOES:
            disciplina = SECOES[chave]

        if (
            paragrafo.text.strip()
            and obter_numeracao(paragrafo) == (3, 0)
        ):
            inicios.append(indice)
            disciplinas[indice] = disciplina

    questoes = []

    for posicao, inicio in enumerate(inicios):
        fim = (
            inicios[posicao + 1]
            if posicao + 1 < len(inicios)
            else len(paragrafos)
        )

        resposta = next(
            (
                indice
                for indice in range(inicio, fim)
                if eh_resolucao(
                    paragrafos[indice].text
                )
                and obter_numeracao(
                    paragrafos[indice]
                ) not in {(3, 1), (11, 0)}
            ),
            None,
        )

        if resposta is None:
            continue

        alternativas = [
            indice
            for indice in range(inicio + 1, resposta)
            if (
                paragrafos[indice].text.strip()
                and obter_numeracao(
                    paragrafos[indice]
                ) in {(3, 1), (11, 0)}
            )
        ]

        if len(alternativas) < 4:
            continue

        alternativas = alternativas[-4:]

        enunciado = "\n".join(
            paragrafos[indice].text.strip()
            for indice in range(
                inicio,
                alternativas[0],
            )
            if paragrafos[indice].text.strip()
        )

        numero = posicao + 1
        gabarito = gabaritos.get(numero)

        if (
            not enunciado
            or gabarito not in {"A", "B", "C", "D"}
        ):
            continue

        resolucao = "\n".join(
            paragrafos[indice].text.strip()
            for indice in range(resposta + 1, fim)
            if paragrafos[indice].text.strip()
        )

        alternativas_html = {}

        for indice_letra, letra in enumerate("ABCD"):
            comeco = alternativas[indice_letra]

            termino = (
                alternativas[indice_letra + 1]
                if indice_letra < 3
                else resposta
            )

            alternativas_html[letra] = html_intervalo(
                paragrafos,
                comeco,
                termino,
                documento,
                salvar,
                prefixo,
            )

        questoes.append({
            "numero": numero,
            "disciplina": disciplinas.get(
                inicio,
                "Português",
            ),
            "enunciado": enunciado,
            "enunciado_html": html_intervalo(
                paragrafos,
                inicio,
                alternativas[0],
                documento,
                salvar,
                prefixo,
            ),
            "alternativas": alternativas_html,
            "gabarito": gabarito,
            "resolucao": resolucao,
            "resolucao_html": html_intervalo(
                paragrafos,
                resposta + 1,
                fim,
                documento,
                salvar,
                prefixo,
            ),
            "apoios": textos_referidos(enunciado),
            "imagem_questao": any(
                imagem_ids(paragrafos[indice])
                for indice in range(
                    inicio,
                    resposta,
                )
            ),
        })

    apoios = extrair_apoios(
        documento,
        inicios[0] if inicios else 0,
        salvar,
        prefixo,
    )

    return {
        "titulo": titulo,
        "subtitulo": tema[:50],
        "ano": ano,
        "questoes": questoes,
        "lidas": len(inicios),
        "gabaritos": len(gabaritos),
        "apoios": apoios,
    }


def extrair_questoes(caminho, salvar=False):
    documento = Document(str(caminho))
    linhas = ler_paragrafos(documento)
    gabaritos = extrair_gabarito(linhas)

    correspondencia = PADRAO_ARQUIVO.match(
        caminho.name
    )

    if not correspondencia:
        raise ValueError(
            f"Nome de arquivo não reconhecido: "
            f"{caminho.name}"
        )

    parte = correspondencia.group(1).upper()
    ano = int(correspondencia.group(2))
    tema = correspondencia.group(3).strip()

    titulo = f"{parte} Simuladão {ano}"[:50]
    prefixo = f"{ano}_{parte}"

    if ano == 2020:
        dados = extrair_questoes_2020(
            documento,
            gabaritos,
            titulo,
            tema,
            ano,
            salvar,
            prefixo,
        )

        dados["arquivo"] = caminho.name
        return dados

    questoes = []
    atual = None
    disciplina = "Português"
    proximo = 1
    resolvendo = False
    primeiro_indice = None
    ultimo_campo = None

    def finalizar():
        nonlocal atual

        if (
            atual
            and atual["enunciado"].strip()
            and all(
                atual["alternativas"][letra].strip()
                for letra in "ABCD"
            )
            and atual["gabarito"] in {
                "A", "B", "C", "D"
            }
        ):
            atual["apoios"] = textos_referidos(
                atual["enunciado"]
            )

            questoes.append(atual)

        atual = None

    for indice, paragrafo in enumerate(
        documento.paragraphs
    ):
        texto = paragrafo.text.strip()
        chave = normalizar(texto).rstrip(" :")

        if chave in SECOES:
            finalizar()
            disciplina = SECOES[chave]
            resolvendo = False
            ultimo_campo = None
            continue

        if chave in {
            "GABARITO",
            "GABARITO OFICIAL",
            "RESPOSTAS",
        }:
            finalizar()
            break

        inicio = (
            PADRAO_QUESTAO.match(texto)
            if texto
            else None
        )

        if (
            inicio
            and int(inicio.group(1)) == proximo
            and 1 <= proximo <= 40
        ):
            finalizar()

            if primeiro_indice is None:
                primeiro_indice = indice

            atual = {
                "numero": proximo,
                "disciplina": disciplina,
                "enunciado": inicio.group(2).strip(),
                "enunciado_html": html_paragrafo(
                    paragrafo,
                    documento,
                    salvar,
                    prefixo,
                ),
                "alternativas": dict.fromkeys(
                    "ABCD",
                    "",
                ),
                "gabarito": gabaritos.get(proximo),
                "resolucao": "",
                "resolucao_html": "",
                "imagem_questao": bool(
                    imagem_ids(paragrafo)
                ),
            }

            proximo += 1
            resolvendo = False
            ultimo_campo = "enunciado"
            continue

        if atual is None:
            continue

        if eh_resolucao(texto):
            resolvendo = True
            ultimo_campo = "resolucao"

            conteudo = html_paragrafo(
                paragrafo,
                documento,
                salvar,
                prefixo,
            )

            atual["resolucao"] += "\n" + texto
            atual["resolucao_html"] += (
                "\n" + conteudo
            )
            continue

        alternativa = (
            PADRAO_ALTERNATIVA.match(texto)
            if texto
            else None
        )

        if alternativa and not resolvendo:
            letra = alternativa.group(1).upper()

            atual["alternativas"][letra] = (
                html_paragrafo(
                    paragrafo,
                    documento,
                    salvar,
                    prefixo,
                )
            )

            ultimo_campo = letra

            atual["imagem_questao"] |= bool(
                imagem_ids(paragrafo)
            )
            continue

        conteudo = html_paragrafo(
            paragrafo,
            documento,
            salvar,
            prefixo,
        )

        if not conteudo:
            continue

        if resolvendo:
            atual["resolucao"] += "\n" + texto
            atual["resolucao_html"] += (
                "\n" + conteudo
            )

        elif ultimo_campo in {
            "A", "B", "C", "D"
        }:
            atual["alternativas"][
                ultimo_campo
            ] += "\n" + conteudo

            atual["imagem_questao"] |= bool(
                imagem_ids(paragrafo)
            )

        elif not any(
            atual["alternativas"].values()
        ):
            atual["enunciado"] += "\n" + texto
            atual["enunciado_html"] += (
                "\n" + conteudo
            )

            atual["imagem_questao"] |= bool(
                imagem_ids(paragrafo)
            )

    finalizar()

    apoios = extrair_apoios(
        documento,
        primeiro_indice or 0,
        salvar,
        prefixo,
    )

    return {
        "titulo": titulo,
        "subtitulo": tema[:50],
        "ano": ano,
        "questoes": questoes,
        "lidas": proximo - 1,
        "gabaritos": len(gabaritos),
        "apoios": apoios,
        "arquivo": caminho.name,
    }


class Command(BaseCommand):
    help = (
        "Analisa e importa os simulados Meta IFRN "
        "com imagens e textos de apoio."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--pasta",
            default=str(PASTA_PADRAO),
        )

        parser.add_argument(
            "--confirmar",
            action="store_true",
        )

    def handle(self, *args, **options):
        pasta = Path(options["pasta"])

        if not pasta.is_dir():
            raise CommandError(
                f"Pasta não encontrada: {pasta}"
            )

        arquivos = sorted(
            caminho
            for caminho in pasta.rglob("*.docx")
            if not caminho.name.startswith("~$")
        )

        if not arquivos:
            raise CommandError(
                "Nenhum DOCX encontrado."
            )

        resultados = []

        self.stdout.write(
            f"Documentos encontrados: {len(arquivos)}"
        )

        for arquivo in arquivos:
            try:
                dados = extrair_questoes(arquivo)
                resultados.append(dados)

                imagens = sum(
                    bool(questao["imagem_questao"])
                    for questao in dados["questoes"]
                )

                alternativas_longas = [
                    (
                        questao["numero"],
                        letra,
                        len(valor),
                    )
                    for questao in dados["questoes"]
                    for letra, valor
                    in questao["alternativas"].items()
                    if len(valor) > 500
                ]

                self.stdout.write(
                    f"\n{arquivo.name}\n"
                    f"  Questões: {len(dados['questoes'])}\n"
                    f"  Gabaritos: {dados['gabaritos']}\n"
                    f"  Textos de apoio: {len(dados['apoios'])}\n"
                    f"  Questões com imagem própria: {imagens}\n"
                    f"  Alternativas acima de 500: "
                    f"{len(alternativas_longas)}"
                )

                for numero, letra, tamanho in alternativas_longas:
                    self.stdout.write(
                        f"    Questão {numero}, "
                        f"alternativa {letra}: "
                        f"{tamanho} caracteres"
                    )

            except Exception as erro:
                self.stdout.write(
                    self.style.ERROR(
                        f"Erro em {arquivo.name}: {erro}"
                    )
                )

        total = sum(
            len(dados["questoes"])
            for dados in resultados
        )

        self.stdout.write(
            f"\nSimulados analisados: {len(resultados)}\n"
            f"Questões identificadas: {total}"
        )

        if not options["confirmar"]:
            self.stdout.write(
                self.style.WARNING(
                    "SIMULAÇÃO: banco não alterado."
                )
            )
            return

        if (
            len(resultados) != 9
            or any(
                len(dados["questoes"]) != 40
                or dados["gabaritos"] < 40
                for dados in resultados
            )
        ):
            raise CommandError(
                "Importação cancelada: são necessários "
                "9 simulados com 40 questões e "
                "40 gabaritos cada."
            )

        if any(
            not dados["apoios"]
            for dados in resultados
        ):
            raise CommandError(
                "Importação cancelada: existem "
                "simulados sem textos de apoio."
            )

        for dados in resultados:
            disponiveis = {
                apoio["numero"]
                for apoio in dados["apoios"]
            }

            referidos = set()

            for questao in dados["questoes"]:
                referidos.update(
                    questao["apoios"]
                )

            faltantes = referidos - disponiveis

            if faltantes:
                raise CommandError(
                    f"Textos de apoio ausentes em "
                    f"{dados['arquivo']}: "
                    f"{sorted(faltantes)}"
                )

        for dados in resultados:
            for questao in dados["questoes"]:
                for letra, valor in (
                    questao["alternativas"].items()
                ):
                    if len(valor) > 500:
                        raise CommandError(
                            f"Alternativa {letra} "
                            f"da questão {questao['numero']} "
                            f"em {dados['arquivo']} "
                            f"excede 500 caracteres."
                        )

        criadas_avaliacoes = 0
        criadas_questoes = 0
        criados_apoios = 0

        with transaction.atomic():
            fonte, _ = Fonte.objects.get_or_create(
                nome="Meta IFRN"
            )

            for dados in resultados:
                existente = Avaliacao.objects.filter(
                    titulo=dados["titulo"],
                    ano=dados["ano"],
                    fonte=fonte,
                    tipo="Simulado",
                ).exists()

                if existente:
                    self.stdout.write(
                        f"Já existe: {dados['titulo']}; "
                        f"ignorado."
                    )
                    continue

                arquivo = next(
                    caminho
                    for caminho in arquivos
                    if caminho.name == dados["arquivo"]
                )

                dados = extrair_questoes(
                    arquivo,
                    salvar=True,
                )

                avaliacao = Avaliacao.objects.create(
                    titulo=dados["titulo"],
                    subtitulo=dados["subtitulo"],
                    fonte=fonte,
                    ano=dados["ano"],
                    tipo="Simulado",
                )

                questoes = []
                mapa_apoios = {}

                for apoio in dados["apoios"]:
                    objeto = TextoApoio.objects.create(
                        titulo=(
                            f"{dados['titulo']} - "
                            f"{apoio['titulo']}"
                        )[:50],
                        texto=apoio["texto"],
                    )

                    mapa_apoios[
                        apoio["numero"]
                    ] = objeto

                    criados_apoios += 1

                for item in dados["questoes"]:
                    disciplina, _ = (
                        Disciplina.objects.get_or_create(
                            nome=item["disciplina"][:20]
                        )
                    )

                    assunto, _ = (
                        Assunto.objects.get_or_create(
                            nome="Simulados Meta IFRN",
                            disciplina=disciplina,
                        )
                    )

                    questao = Questao.objects.create(
                        ano=dados["ano"],
                        disciplina=disciplina,
                        assunto=assunto,
                        fonte=fonte,
                        dificuldade="Média",
                        enunciado=item["enunciado_html"],
                        gabarito_comentado=(
                            item["resolucao_html"]
                        ),
                        alternativa_a=(
                            item["alternativas"]["A"]
                        ),
                        alternativa_b=(
                            item["alternativas"]["B"]
                        ),
                        alternativa_c=(
                            item["alternativas"]["C"]
                        ),
                        alternativa_d=(
                            item["alternativas"]["D"]
                        ),
                        alternativa_correta=(
                            item["gabarito"]
                        ),
                    )

                    for numero in item["apoios"]:
                        mapa_apoios[
                            numero
                        ].questoes.add(questao)

                    questoes.append(questao)
                    criadas_questoes += 1

                avaliacao.questoes.set(questoes)
                criadas_avaliacoes += 1

        self.stdout.write(
            self.style.SUCCESS(
                "Importação concluída. "
                f"Simulados: {criadas_avaliacoes}; "
                f"questões: {criadas_questoes}; "
                f"textos de apoio: {criados_apoios}."
            )
        )
